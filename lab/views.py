import csv
import io
import json
from zipfile import BadZipFile
from itertools import islice
from datetime import timedelta
from urllib.parse import urlparse
from django.conf import settings
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.password_validation import validate_password
from django.contrib.auth.tokens import default_token_generator
from django.core.exceptions import ValidationError
from django.core.mail import send_mail
from django.core.validators import validate_email
from django.db import IntegrityError, transaction
from django.db.models import F
from django.http import JsonResponse
from django.middleware.csrf import get_token
from django.utils import timezone
from django.utils.dateparse import parse_datetime
from django.utils.encoding import force_bytes
from django.utils.http import urlsafe_base64_encode, urlsafe_base64_decode
from django.views.decorators.cache import never_cache
from openpyxl import load_workbook
from .models import User, Attendance, Visit, Item, Loan, Event, Notification, Mutation, MailTask

def lockout(request, credentials=None, *args, **kwargs):
    return JsonResponse({'error':'Demasiados intentos. Intenta nuevamente en 15 minutos.'}, status=429)

def require(condition, message, status=400):
    if not condition:
        raise ApiError(message, status)

class ApiError(Exception):
    def __init__(self, message, status=400):
        self.message, self.status = message, status

def text(data, key, length=300):
    value = str(data.get(key, '')).strip()
    require(0 < len(value) <= length, f'Revisa el campo {key}.')
    return value

def number(data, key, minimum=1):
    value = data.get(key)
    require(not isinstance(value, bool) and str(value).isdigit(), f'Revisa la cantidad: {key}.')
    result = int(value)
    require(minimum <= result <= 1000000, f'Cantidad fuera de rango: {key}.')
    return result

def person(u):
    return {'id':u.id,'name':u.first_name,'dni':u.username,'email':u.email,'role':u.role,'active':u.is_active}

def event(user, message, kind):
    return Event.objects.create(actor=user, text=message, kind=kind)

def send_access(task):
    user = task.user
    link = f'{settings.PUBLIC_URL}/?uid={urlsafe_base64_encode(force_bytes(user.pk))}&token={default_token_generator.make_token(user)}'
    send_mail('Tu acceso a NonHuman', f'Crea o restablece tu contrasena en:\n{link}\n\nTu usuario es tu DNI. El enlace vence en una hora y se puede usar una sola vez.\nSi no solicitaste este correo, ignoralo.\n\nNonHuman', settings.DEFAULT_FROM_EMAIL, [user.email])
    task.sent = timezone.now()
    task.attempts += 1
    task.save()

def dispatch_mail(user_id):
    task = MailTask.objects.filter(user_id=user_id, sent=None).order_by('created').first()
    if task:
        try:
            send_access(task)
        except Exception:
            task.attempts += 1
            task.save(update_fields=['attempts'])

def parse_people(rows, allow_existing=False):
    require(isinstance(rows,list) and len(rows) <= 500, 'Importa hasta 500 personas a la vez.')
    result, seen_dni, seen_email = [], set(), set()
    for row in rows:
        require(isinstance(row,dict), 'Fila invalida.')
        row = {str(k).strip().lower():v for k,v in row.items()}
        raw_dni = row.get('dni', '')
        if isinstance(raw_dni,float) and raw_dni.is_integer():
            raw_dni = int(raw_dni)
        dni = str(raw_dni).strip()
        email = str(row.get('correo', row.get('email', ''))).strip().lower()
        name = str(row.get('nombre', row.get('name', '')) or '').strip()
        if not name:
            name = ' '.join(str(row.get(k) or '').strip() for k in ('nombres','apellidos')).strip()
        name = ' '.join(name.split())
        role = str(row.get('rol', row.get('role', 'member')) or 'member').strip().lower()
        role = {'administrador':'admin','miembro':'member','externo':'external'}.get(role, role)
        error = ''
        try:
            require(len(dni) == 8 and dni.isdigit(), 'DNI: 8 digitos.')
            require(0 < len(name) <= 100, 'Nombre requerido (hasta 100 caracteres).')
            validate_email(email)
            require(role in ('admin','member','external'), 'Rol invalido.')
            require(dni not in seen_dni and (allow_existing or not User.objects.filter(username=dni).exists()), 'DNI duplicado.')
            require(email not in seen_email and (allow_existing or not User.objects.filter(email__iexact=email).exists()), 'Correo duplicado.')
        except (ApiError, ValidationError) as exc:
            error = exc.message if isinstance(exc, ApiError) else 'Correo invalido.'
        seen_dni.add(dni)
        seen_email.add(email)
        result.append({'dni':dni,'email':email,'name':name,'role':role,'error':error})
    return result

def read_people_file(file, allow_existing=False):
    if file.name.lower().endswith('.csv'):
        rows = list(islice(csv.DictReader(io.StringIO(file.read().decode('utf-8-sig'))), 501))
    else:
        require(file.name.lower().endswith('.xlsx'), 'Usa XLSX o CSV.')
        book = load_workbook(file,read_only=True,data_only=True)
        try:
            iterator = book.active.iter_rows(values_only=True)
            headers = [str(h or '').strip().lower() for h in next(iterator)]
            rows = list(islice((dict(zip(headers,row)) for row in iterator if any(v is not None for v in row)),501))
        finally:
            book.close()
    return parse_people(rows, allow_existing=allow_existing)

def snapshot(user):
    # shortcut: full histories suit this small lab; paginate when payloads grow.
    admin = user.role == 'admin'
    attendance = Attendance.objects.select_related('user').order_by('-entered')
    if not admin:
        attendance = attendance.filter(user=user)
    payload = {'user':person(user),'demo':settings.DEMO_MODE,'attendance':[
        {'id':a.id,'user_id':a.user_id,'name':a.user.first_name,'role':a.user.role,'entered':a.entered,'exited':a.exited}
        for a in attendance
    ],'inside':Attendance.objects.filter(user=user,exited=None).exists()}
    if user.role != 'external':
        payload['items'] = list(Item.objects.filter(active=True).order_by('name').values())
        loans = Loan.objects.select_related('user','item').order_by('-created')
        visits = Visit.objects.select_related('host').order_by('-entered')
        if not admin:
            loans = loans.filter(user=user)
            visits = visits.filter(host=user)
        payload['loans'] = [{'id':l.id,'user_id':l.user_id,'name':l.user.first_name,'item_id':l.item_id,'item':l.item.name,'quantity':l.quantity,'returned':l.returned,'reason':l.reason,'due':l.due,'created':l.created,'consumable':l.item.consumable} for l in loans]
        payload['visits'] = [{'id':v.id,'host_id':v.host_id,'host':v.host.first_name,'name':v.name,'reason':v.reason,'entered':v.entered,'exited':v.exited} for v in visits]
    if admin:
        payload['people'] = [person(u) for u in User.objects.order_by('first_name')]
        payload['events'] = [{'id':e.id,'name':e.actor.first_name,'text':e.text,'kind':e.kind,'created':e.created} for e in Event.objects.select_related('actor').order_by('-created')[:100]]
        payload['notifications'] = [{'id':n.id,'text':n.event.text,'name':n.event.actor.first_name,'created':n.event.created,'read':n.read} for n in Notification.objects.filter(recipient=user).select_related('event__actor').order_by('-id')[:100]]
        payload['mail_pending'] = MailTask.objects.filter(sent=None).count()
    return payload

@never_cache
def api(request, action):
    try:
        if request.method == 'GET':
            if action == 'session':
                return JsonResponse({'csrf':get_token(request),'user':person(request.user) if request.user.is_authenticated else None,'demo':settings.DEMO_MODE})
            require(request.user.is_authenticated, 'Inicia sesion para continuar.', 401)
            require(action == 'state', 'No encontrado.', 404)
            return JsonResponse(snapshot(request.user))
        require(request.method == 'POST', 'Metodo no permitido.', 405)
        data = json.loads(request.body or '{}') if request.content_type == 'application/json' else {}
        require(isinstance(data, dict), 'Solicitud invalida.')
        if action == 'login':
            user = authenticate(request, username=str(data.get('dni','')), password=str(data.get('password','')))
            require(user is not None, 'DNI o contrasena incorrectos.', 401)
            login(request, user)
            return JsonResponse({'user':person(user),'csrf':get_token(request)})
        if action == 'logout':
            logout(request)
            return JsonResponse({'ok':True})
        if action == 'recover':
            email = str(data.get('email','')).strip().lower()
            user = User.objects.filter(email__iexact=email,is_active=True).first()
            recent = timezone.now() - timedelta(minutes=15)
            if user and not MailTask.objects.filter(user=user,created__gt=recent).exists() and MailTask.objects.filter(created__gt=recent).count() < 30:
                MailTask.objects.create(user=user)
                dispatch_mail(user.id)
            return JsonResponse({'ok':True})
        if action == 'set-password':
            try:
                user = User.objects.get(pk=urlsafe_base64_decode(data.get('uid','')).decode(),is_active=True)
            except (ValueError, UnicodeDecodeError, User.DoesNotExist, OverflowError):
                raise ApiError('Enlace invalido o vencido.')
            with transaction.atomic():
                user = User.objects.select_for_update().get(pk=user.pk)
                require(default_token_generator.check_token(user, str(data.get('token',''))), 'Enlace invalido o vencido.')
                password = text(data,'password',128)
                validate_password(password,user)
                user.set_password(password)
                user.save()
            return JsonResponse({'ok':True})
        require(request.user.is_authenticated, 'Inicia sesion para continuar.',401)
        user = request.user
        if action == 'import-preview':
            require(user.role == 'admin','No autorizado.',403)
            file = request.FILES.get('file')
            require(file is not None and file.size <= 2*1024*1024, 'Selecciona un archivo XLSX o CSV de hasta 2 MB.')
            return JsonResponse({'rows':read_people_file(file)})
        with transaction.atomic():
            user = User.objects.select_for_update().get(pk=user.pk)
            require(user.is_active,'Cuenta desactivada.',403)
            key = request.headers.get('Idempotency-Key','')
            require(0 < len(key) <= 64, 'Falta identificador de operacion.')
            if Mutation.objects.filter(user=user,key=key).exists():
                return JsonResponse({'ok':True,'replayed':True})
            Mutation.objects.create(user=user,key=key)
            result = mutate(user,action,data)
        return JsonResponse({'ok':True,**(result or {})})
    except ApiError as exc:
        return JsonResponse({'error':exc.message},status=exc.status)
    except ValidationError as exc:
        return JsonResponse({'error':' '.join(exc.messages)},status=400)
    except IntegrityError:
        return JsonResponse({'error':'El registro ya existe o cambio. Actualiza e intenta nuevamente.'},status=409)
    except (ValueError,TypeError,KeyError,StopIteration,BadZipFile):
        return JsonResponse({'error':'Datos invalidos. Revisa el formulario.'},status=400)
    except (User.DoesNotExist,Item.DoesNotExist,Loan.DoesNotExist,Visit.DoesNotExist,Attendance.DoesNotExist):
        return JsonResponse({'error':'Registro no encontrado.'},status=404)

def mutate(user, action, data):
    now = timezone.now()
    if action == 'check-in':
        require(not Attendance.objects.filter(user=user,exited=None).exists(),'Ya registraste tu entrada.')
        Attendance.objects.create(user=user,entered=now)
        event(user,'Registro de entrada','attendance')
        return
    if action == 'check-out':
        require(Attendance.objects.filter(user=user,exited=None).update(exited=now) == 1,'No tienes una entrada abierta.')
        event(user,'Registro de salida','attendance')
        return {'open_visits':Visit.objects.filter(host=user,exited=None).count()}
    require(user.role != 'external','Tu cuenta solo permite asistencia e historial.',403)
    if action == 'visit-create':
        Visit.objects.create(host=user,name=text(data,'name',100),reason=text(data,'reason'),entered=now)
        event(user,f'Ingreso de visitante: {data["name"]}','visit')
        return
    if action == 'visit-close':
        visit = Visit.objects.select_for_update().get(pk=number(data,'id'))
        require(visit.host_id == user.id or user.role == 'admin','No autorizado.',403)
        require(visit.exited is None,'La visita ya finalizo.')
        visit.exited = now
        visit.save()
        event(user,f'Salida de visitante: {visit.name}','visit')
        return
    if action == 'checkout':
        reason = text(data,'reason')
        lines = data.get('items',[])
        require(isinstance(lines,list) and 0 < len(lines) <= 50,'Selecciona articulos.')
        require(all(isinstance(line,dict) for line in lines),'Articulos invalidos.')
        ids = [number(line,'id') for line in lines]
        require(len(ids) == len(set(ids)), 'Articulos duplicados.')
        due = parse_datetime(str(data.get('due','')))
        if due and timezone.is_naive(due):
            due = timezone.make_aware(due)
        descriptions = []
        for line in sorted(lines,key=lambda line:int(line['id'])):
            item = Item.objects.select_for_update().get(pk=line['id'],active=True)
            quantity = number(line,'quantity')
            require(item.available >= quantity,f'Stock insuficiente: {item.name}.')
            require(item.consumable or (due is not None and due > now),'Indica una devolucion futura.')
            updated = Item.objects.filter(pk=item.pk,available__gte=quantity).update(available=F('available')-quantity, total=F('total')-quantity if item.consumable else F('total'))
            require(updated == 1,'La disponibilidad cambio. Intenta nuevamente.',409)
            Loan.objects.create(user=user,item=item,quantity=quantity,reason=reason,due=None if item.consumable else due,created=now)
            descriptions.append(f'{quantity} × {item.name}')
        e = event(user,('Retiro: '+', '.join(descriptions))[:400],'checkout')
        Notification.objects.bulk_create([Notification(event=e,recipient=admin) for admin in User.objects.filter(role='admin',is_active=True)])
        return
    if action == 'return':
        loan = Loan.objects.select_for_update().select_related('item').get(pk=number(data,'id'))
        require(loan.user_id == user.id, 'Solo puedes devolver tus propios prestamos.',403)
        require(not loan.item.consumable,'Un consumo no es un prestamo.')
        quantity = number(data,'quantity')
        require(quantity <= loan.quantity-loan.returned,'Cantidad superior a lo pendiente.')
        damaged = data.get('damaged') is True
        loan.returned += quantity
        loan.save()
        item = Item.objects.select_for_update().get(pk=loan.item_id)
        if damaged:
            Item.objects.filter(pk=item.pk).update(unavailable=F('unavailable')+quantity)
        else:
            Item.objects.filter(pk=item.pk).update(available=F('available')+quantity)
        event(user,f'Devolucion: {quantity} × {item.name}'+(' (averiado)' if damaged else ''),'return')
        return
    require(user.role == 'admin','Solo administradores.',403)
    if action == 'notifications-read':
        Notification.objects.filter(recipient=user).update(read=True)
    elif action in ('person-create','import-people'):
        rows = parse_people(data.get('rows',[]) if action == 'import-people' else [data])
        require(rows and all(not row['error'] for row in rows),'Hay filas invalidas o duplicadas.')
        for row in rows:
            new = User.objects.create_user(username=row['dni'],email=row['email'],first_name=row['name'],role=row['role'],password=None)
            MailTask.objects.create(user=new)
            event(user,f'Alta de cuenta: {new.first_name}','people')
            transaction.on_commit(lambda uid=new.id: dispatch_mail(uid))
    elif action == 'person-toggle':
        admins = list(User.objects.select_for_update().filter(role='admin',is_active=True).order_by('id'))
        target = User.objects.select_for_update().get(pk=number(data,'id'))
        require(target.id != user.id,'No puedes desactivar tu propia cuenta.')
        require(not (target.role == 'admin' and target.is_active and len(admins) <= 1),'Debe existir un administrador activo.')
        target.is_active = not target.is_active
        if not target.is_active:
            # Updating the password hash revokes old sessions even after reactivation.
            target.set_unusable_password()
        target.save()
        if target.is_active:
            MailTask.objects.create(user=target)
            transaction.on_commit(lambda:dispatch_mail(target.id))
        event(user,f'{"Activacion" if target.is_active else "Baja"}: {target.first_name}','people')
    elif action == 'invite':
        target = User.objects.get(pk=number(data,'id'),is_active=True)
        require(not MailTask.objects.filter(user=target,created__gt=now-timedelta(minutes=2)).exists(),'Espera dos minutos antes de reenviar.')
        MailTask.objects.create(user=target)
        transaction.on_commit(lambda:dispatch_mail(target.id))
    elif action == 'item-create':
        image = str(data.get('image',''))
        require(not image or image.startswith('/assets/') or urlparse(image).scheme == 'https','Imagen: usa una URL HTTPS.')
        quantity = number(data,'quantity',0)
        Item.objects.create(code=text(data,'code',30),name=text(data,'name',120),category=text(data,'category',50),location=text(data,'location',100),total=quantity,available=quantity,consumable=data.get('consumable') is True,minimum=number(data,'minimum',0),image=image[:500])
        event(user,f'Alta de articulo: {data["name"]}','stock')
    elif action == 'stock':
        item = Item.objects.select_for_update().get(pk=number(data,'id'))
        quantity = number(data,'quantity')
        reason = text(data,'reason')
        operation = data.get('operation','add')
        require(operation in ('add','remove','repair'),'Operacion invalida.')
        if operation == 'add':
            item.total += quantity
            item.available += quantity
        elif operation == 'remove':
            require(item.available >= quantity,'Solo se puede dar de baja stock disponible.')
            item.total -= quantity
            item.available -= quantity
        else:
            require(item.unavailable >= quantity,'No hay suficientes unidades averiadas.')
            item.unavailable -= quantity
            item.available += quantity
        item.save()
        event(user,f'{operation}: {quantity} × {item.name}. {reason}'[:400],'stock')
    else:
        raise ApiError('Accion no encontrada.',404)
